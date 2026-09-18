"""Report generation utilities for ThermoShelter360.

Provides programmatic export of simulation results, thermal balance logs,
and material comparison tables to formatted Excel (.xlsx) and PDF (.pdf) documents.
"""

import io
from typing import Any, Dict, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def generate_excel_report(
    summary_data: Dict[str, Any],
    timeseries_df: pd.DataFrame,
    comparison_df: Optional[pd.DataFrame] = None,
) -> bytes:
    """Generate a multi-tab Excel workbook containing complete simulation data.

    Tabs:
        1. "Executive Summary" - Project metadata, KPI metrics, comfort stats
        2. "Hourly Simulation" - Complete 24-hr/dynamic thermal timeseries
        3. "Material Comparison" - Comparative metrics across wall assemblies

    Args:
        summary_data: Key-value dictionary of metadata and KPI results.
        timeseries_df: DataFrame of simulation time series.
        comparison_df: Optional DataFrame of material comparison runs.

    Returns:
        Bytes buffer of the generated Excel workbook.
    """
    wb = openpyxl.Workbook()

    # Define common styles
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # -------------------------------------------------------------
    # Sheet 1: Executive Summary
    # -------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary.cell(row=1, column=1, value="ThermoShelter360 — Thermal Simulation Report").font = title_font
    ws_summary.cell(row=2, column=1, value="High-Altitude Cold Region Passive Shelter Analysis").font = Font(name="Calibri", size=11, italic=True, color="595959")

    ws_summary.cell(row=4, column=1, value="Metric / Parameter").font = header_font
    ws_summary.cell(row=4, column=1).fill = header_fill
    ws_summary.cell(row=4, column=2, value="Value").font = header_font
    ws_summary.cell(row=4, column=2).fill = header_fill

    current_row = 5
    for k, v in summary_data.items():
        cell_k = ws_summary.cell(row=current_row, column=1, value=str(k))
        cell_v = ws_summary.cell(row=current_row, column=2, value=str(v))
        cell_k.font = bold_font
        cell_v.font = regular_font
        cell_k.border = thin_border
        cell_v.border = thin_border
        current_row += 1

    ws_summary.column_dimensions["A"].width = 32
    ws_summary.column_dimensions["B"].width = 30

    # -------------------------------------------------------------
    # Sheet 2: Hourly Simulation
    # -------------------------------------------------------------
    ws_sim = wb.create_sheet(title="Hourly Simulation")
    ws_sim.views.sheetView[0].showGridLines = True

    # Write headers
    for col_idx, col_name in enumerate(timeseries_df.columns, start=1):
        cell = ws_sim.cell(row=1, column=col_idx, value=str(col_name).replace("_", " ").title())
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Write data rows
    for row_idx, row in timeseries_df.iterrows():
        for col_idx, value in enumerate(row, start=1):
            cell = ws_sim.cell(row=row_idx + 2, column=col_idx, value=value)
            cell.font = regular_font
            cell.border = thin_border
            if isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal="right")

    for col in ws_sim.columns:
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_sim.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # -------------------------------------------------------------
    # Sheet 3: Material Comparison
    # -------------------------------------------------------------
    if comparison_df is not None and not comparison_df.empty:
        ws_comp = wb.create_sheet(title="Material Comparison")
        ws_comp.views.sheetView[0].showGridLines = True

        for col_idx, col_name in enumerate(comparison_df.columns, start=1):
            cell = ws_comp.cell(row=1, column=col_idx, value=str(col_name).replace("_", " ").title())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, row in comparison_df.iterrows():
            for col_idx, value in enumerate(row, start=1):
                cell = ws_comp.cell(row=row_idx + 2, column=col_idx, value=value)
                cell.font = regular_font
                cell.border = thin_border
                if isinstance(value, (int, float)):
                    cell.alignment = Alignment(horizontal="right")

        for col in ws_comp.columns:
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws_comp.column_dimensions[col_letter].width = max(max_len + 3, 16)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_pdf_report(
    project_name: str,
    location: str,
    summary_data: Dict[str, Any],
    comfort_metrics: Dict[str, Any],
    geometry: Dict[str, Any],
    material: Dict[str, Any],
) -> bytes:
    """Generate a clean executive PDF report for the thermal simulation.

    Args:
        project_name: Name of the shelter project.
        location: Geographic region.
        summary_data: Dictionary of summary KPI outputs.
        comfort_metrics: Thermal comfort breakdown metrics.
        geometry: Geometry inputs dictionary.
        material: Wall material specification.

    Returns:
        Bytes buffer of the generated PDF file.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    primary_color = colors.HexColor("#1A365D")
    secondary_color = colors.HexColor("#2B6CB0")
    light_bg = colors.HexColor("#F7FAFC")
    border_color = colors.HexColor("#E2E8F0")

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontSize=20,
        leading=24,
        textColor=primary_color,
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#718096"),
        spaceAfter=12,
    )
    heading_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontSize=12,
        leading=16,
        textColor=secondary_color,
        spaceBefore=8,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "BodyTextCustom",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#2D3748"),
    )
    disclaimer_style = ParagraphStyle(
        "DisclaimerText",
        parent=styles["Italic"],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#A0AEC0"),
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph("ThermoShelter360 — Engineering Thermal Report", title_style))
    story.append(
        Paragraph(
            f"Project: <b>{project_name}</b> | Location: <b>{location}</b> | Passive Solar Thermal Simulation",
            subtitle_style,
        )
    )
    story.append(HRFlowable(width="100%", thickness=1.5, color=secondary_color, spaceAfter=10))

    # KPI Summary Table
    story.append(Paragraph("Key Performance Indicators (Simulation Results)", heading_style))
    kpi_rows = [["Metric", "Result", "Thermal Comfort Evaluation", "Value"]]
    kpi_rows.append(["Average Indoor Temp", f"{summary_data.get('avg_indoor_temp', 'N/A')} °C", "Hours in Comfort Range (18-24°C)", f"{comfort_metrics.get('hours_in_comfort_range', 'N/A')} hrs"])
    kpi_rows.append(["Minimum Indoor Temp", f"{summary_data.get('min_indoor_temp', 'N/A')} °C", "Percent Comfortable Time", f"{comfort_metrics.get('percent_comfortable', 'N/A')} %"])
    kpi_rows.append(["Maximum Indoor Temp", f"{summary_data.get('max_indoor_temp', 'N/A')} °C", "Supplemental Heating Demand", f"{summary_data.get('heating_requirement', 'N/A')} kWh"])
    kpi_rows.append(["Total Envelope Heat Loss", f"{summary_data.get('total_heat_loss', 'N/A')} kWh", "Total Solar Heat Gain", f"{summary_data.get('solar_gain', 'N/A')} kWh"])

    t_kpi = Table(kpi_rows, colWidths=[2.2 * inch, 1.3 * inch, 2.5 * inch, 1.2 * inch])
    t_kpi.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), primary_color),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light_bg]),
                ("GRID", (0, 0), (-1, -1), 0.5, border_color),
                ("ALIGN", (1, 1), (1, -1), "RIGHT"),
                ("ALIGN", (3, 1), (3, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(t_kpi)
    story.append(Spacer(1, 10))

    # Architecture & Construction Parameters Table
    story.append(Paragraph("Shelter Geometry & Envelope Assembly", heading_style))
    mat_name = material.get("name", material.get("material", "Custom Assembly"))
    geom_rows = [
        ["Parameter", "Specification", "Parameter", "Specification"],
        ["Length x Width x Height", f"{geometry.get('length', 6.0)}m x {geometry.get('width', 4.0)}m x {geometry.get('height', 2.8)}m", "Wall Assembly", f"{mat_name}"],
        ["Wall Thickness", f"{geometry.get('wall_thickness', 0.25)} m", "Wall Conductivity (k)", f"{material.get('thermal_conductivity', 0.72)} W/(m·K)"],
        ["Window Area", f"{geometry.get('window_area', 3.0)} m²", "Window U-value", f"{geometry.get('window_u_value', 1.8)} W/(m²·K)"],
        ["Door Area", f"{geometry.get('door_area', 1.8)} m²", "Roof U-value", f"{geometry.get('roof_u_value', 0.35)} W/(m²·K)"],
    ]

    t_geom = Table(geom_rows, colWidths=[2.2 * inch, 1.3 * inch, 2.2 * inch, 1.5 * inch])
    t_geom.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), secondary_color),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light_bg]),
                ("GRID", (0, 0), (-1, -1), 0.5, border_color),
                ("ALIGN", (1, 1), (1, -1), "RIGHT"),
                ("ALIGN", (3, 1), (3, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(t_geom)
    story.append(Spacer(1, 15))

    # Engineering Disclaimer
    story.append(
        Paragraph(
            "<b>Engineering Disclaimer:</b> This is a simplified conceptual thermal model intended for educational "
            "and preliminary design exploration of passive shelters in cold high-altitude environments. It is not a substitute "
            "for certified engineering calculations, on-site measurements, or comprehensive multi-dimensional CFD/thermal simulation.",
            disclaimer_style,
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
